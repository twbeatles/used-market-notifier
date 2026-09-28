"""Regression: exported CSV/XLSX never contains live formulas from listing text (ISSUE-008)."""

import csv
import os
import tempfile
import unittest

from export_manager import ExportManager, sanitize_cell

DANGEROUS = ['=HYPERLINK("http://example.invalid","클릭")', "+cmd", "-2+3", "@SUM(A1)", "\t=1", "\r=1"]


class ExportSanitizeTest(unittest.TestCase):
    def test_sanitize_cell(self):
        for value in DANGEROUS:
            self.assertEqual(sanitize_cell(value), "'" + value)
        self.assertEqual(sanitize_cell("아이폰 15 프로"), "아이폰 15 프로")
        self.assertEqual(sanitize_cell(12000), 12000)
        self.assertIsNone(sanitize_cell(None))

    def test_excel_cells_are_strings(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("openpyxl not installed")
        rows = [{"title": v, "price": 10000} for v in DANGEROUS] + [{"title": "정상 제목", "price": 5}]
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "out.xlsx")
            ok, _ = ExportManager.export_to_excel(rows, path, ["title", "price"])
            self.assertTrue(ok)
            ws = load_workbook(path)["매물 목록"]
            for row in range(2, len(rows) + 2):
                self.assertNotEqual(ws.cell(row=row, column=1).data_type, "f")
            self.assertEqual(ws.cell(row=len(rows) + 1, column=1).value, "정상 제목")
            self.assertEqual(ws.cell(row=2, column=2).value, 10000)

    def test_csv_values_are_prefixed(self):
        rows = [{"title": v} for v in DANGEROUS]
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "out.csv")
            ok, _ = ExportManager.export_to_csv(rows, path, ["title"])
            self.assertTrue(ok)
            with open(path, encoding="utf-8-sig", newline="") as f:
                values = [r["title"] for r in csv.DictReader(f)]
        self.assertTrue(all(v.startswith("'") for v in values))


if __name__ == "__main__":
    unittest.main()
